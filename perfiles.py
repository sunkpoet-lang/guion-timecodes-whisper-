# -*- coding: utf-8 -*-
"""
Perfiles de estilo para traducir: cada traductor (o cada serie) tiene su perfil
con sus reglas en lenguaje natural y su glosario. Se guardan como JSON en
CARPETA_DATOS/perfiles y se pueden exportar/importar para compartirlos.

Perfil:
    {
      "id", "nombre",
      "idioma_origen": "en",                  # código, o "auto"
      "variante": "español latino neutro",    # a qué español se traduce
      "instrucciones": "…",                   # reglas del traductor (p. ej. el estilo.md)
      "acotacion": "(REAC)",                  # si no está vacío, todo paréntesis se reemplaza por esto
      "mayusculas": false,                    # pasar la traducción a MAYÚSCULAS (convención de rayado)
      "glosario": {
        "personajes": [{"original", "doblaje", "variantes": [], "en_dialogo", "tipo", "notas"}],
        "terminos":   [{"original", "doblaje", "notas"}]
      }
    }

El glosario se puede importar de un Excel con las pestañas PERSONAJES y
TERMINOS (el formato de glosario.xlsx por proyecto), o de cualquier hoja con
columnas ORIGINAL y DOBLAJE / TRADUCCIÓN.
"""

import json
import os
import re
import unicodedata
import uuid

from rutas import CARPETA_DATOS

CARPETA = os.path.join(CARPETA_DATOS, "perfiles")
os.makedirs(CARPETA, exist_ok=True)

PERFIL_INICIAL = {
    "id": "latino-neutro",
    "nombre": "Español latino neutro",
    "idioma_origen": "en",
    "variante": "español latino neutro",
    "instrucciones": (
        "- Español latino neutro, entendible en toda Latinoamérica: sin regionalismos marcados ni voseo.\n"
        "- Tuteo entre personajes jóvenes o de confianza; usted solo con desconocidos o figuras de autoridad.\n"
        "- Nada de expresiones de España (vale, tío, coger, vosotros).\n"
        "- Interjecciones y onomatopeyas adaptadas al español (Ay, Oye, Uf) salvo que la regla diga otra cosa."
    ),
    "acotacion": "",
    "mayusculas": False,
    "glosario": {"personajes": [], "terminos": []},
}


def _ruta(id_perfil):
    if not re.fullmatch(r"[\w\-]+", id_perfil or ""):
        raise ValueError("Identificador de perfil inválido.")
    return os.path.join(CARPETA, f"{id_perfil}.json")


def _completar(perfil):
    """Rellena los campos que falten (perfiles de versiones anteriores o importados a mano)."""
    base = json.loads(json.dumps(PERFIL_INICIAL))
    base.update({k: v for k, v in perfil.items() if k != "glosario"})
    glosario = perfil.get("glosario") or {}
    base["glosario"] = {"personajes": glosario.get("personajes") or [], "terminos": glosario.get("terminos") or []}
    return base


def listar():
    perfiles = []
    for nombre in sorted(os.listdir(CARPETA)):
        if nombre.endswith(".json"):
            try:
                with open(os.path.join(CARPETA, nombre), encoding="utf-8") as f:
                    perfiles.append(_completar(json.load(f)))
            except (OSError, ValueError):
                continue
    if not perfiles:
        guardar(PERFIL_INICIAL)
        perfiles = [_completar(PERFIL_INICIAL)]
    return perfiles


def obtener(id_perfil):
    with open(_ruta(id_perfil), encoding="utf-8") as f:
        return _completar(json.load(f))


def guardar(perfil):
    perfil = _completar(perfil)
    if not perfil.get("id"):
        base = re.sub(r"[^a-z0-9]+", "-", _sin_acentos(perfil["nombre"]).lower()).strip("-") or "perfil"
        perfil["id"] = f"{base}-{uuid.uuid4().hex[:4]}"
    with open(_ruta(perfil["id"]), "w", encoding="utf-8") as f:
        json.dump(perfil, f, ensure_ascii=False, indent=2)
    return perfil


def borrar(id_perfil):
    try:
        os.remove(_ruta(id_perfil))
    except FileNotFoundError:
        pass


# ------------------------------------------------------------------
# Importar glosario y estilo
# ------------------------------------------------------------------

def _sin_acentos(texto):
    return unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode("ascii")


def _clave(texto):
    return re.sub(r"[^A-Z]", "", _sin_acentos(str(texto or "")).upper())


def _texto(valor):
    return "" if valor is None else str(valor).strip()


def importar_glosario(ruta_xlsx):
    """Lee un glosario de Excel. Devuelve {"personajes": [...], "terminos": [...]}."""
    from openpyxl import load_workbook

    libro = load_workbook(ruta_xlsx, read_only=True, data_only=True)
    personajes, terminos = [], []

    for hoja in libro.worksheets:
        filas = [f for f in hoja.iter_rows(values_only=True)]
        if not filas:
            continue
        encabezado = [_clave(c) for c in filas[0]]

        def col(*nombres):
            for n in nombres:
                if n in encabezado:
                    return encabezado.index(n)
            return None

        i_orig = col("ORIGINAL", "SOURCE", "INGLES", "ENGLISH")
        i_dob = col("DOBLAJE", "TRADUCCION", "ESPANOL", "SPANISH", "DESTINO")
        if i_orig is None or i_dob is None:
            continue
        i_var, i_dial, i_tipo, i_notas = col("VARIANTES"), col("ENDIALOGO"), col("TIPO"), col("NOTAS", "NOTA")
        es_personajes = _clave(hoja.title) in ("PERSONAJES", "CHARACTERS", "PERSONAJE")

        for fila in filas[1:]:
            fila = list(fila) + [None] * (len(encabezado) - len(fila))
            original, doblaje = _texto(fila[i_orig]), _texto(fila[i_dob])
            if not original or not doblaje:
                continue
            notas = _texto(fila[i_notas]) if i_notas is not None else ""
            if es_personajes:
                variantes = _texto(fila[i_var]) if i_var is not None else ""
                personajes.append({
                    "original": original, "doblaje": doblaje,
                    "variantes": [v.strip() for v in variantes.split(";") if v.strip()],
                    "en_dialogo": _texto(fila[i_dial]) if i_dial is not None else "",
                    "tipo": (_texto(fila[i_tipo]) if i_tipo is not None else "") or "personaje",
                    "notas": notas,
                })
            else:
                tipo = _texto(fila[i_tipo]) if i_tipo is not None else ""
                terminos.append({"original": original, "doblaje": doblaje,
                                 "notas": " · ".join(x for x in (tipo if tipo == "guia" else "", notas) if x)})
    return {"personajes": personajes, "terminos": terminos}


def importar_estilo(ruta):
    """Texto de un estilo.md / .txt para usar como instrucciones del perfil."""
    with open(ruta, encoding="utf-8-sig", errors="replace") as f:
        return f.read().strip()
